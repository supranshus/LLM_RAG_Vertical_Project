"""Read employee records from an Excel workbook."""
from __future__ import annotations

import hashlib
import math
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

def import_employees(content: bytes) -> dict[str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise ValueError("Excel import requires openpyxl. Install the project requirements.") from exc
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("The employee workbook must be 5 MB or smaller.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ValueError("Could not read the file as an Excel workbook.") from exc
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows, ())
        columns = [str(value or "").strip().lower().replace(" ", "_") for value in header]
        if "name" not in columns:
            raise ValueError("The workbook must have a name column in its first row.")
        known = [column for column in columns if column in {"id", "name", "role", "experience_years", "skills"}]
        if len(known) != len(set(known)):
            raise ValueError("The workbook contains duplicate employee columns.")
        employees = {}
        for row_number, values in enumerate(rows, start=2):
            if row_number > 5001:
                raise ValueError("The workbook must contain no more than 5,000 employee rows.")
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            row = dict(zip(columns, values))
            name = str(row.get("name") or "").strip()
            role = str(row.get("role") or "").strip()
            if not name:
                raise ValueError(f"Row {row_number}: employee name is required.")
            experience = row.get("experience_years")
            if experience is not None and str(experience).strip():
                try:
                    experience = float(experience)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"Row {row_number}: experience_years must be a non-negative number.") from exc
                if not math.isfinite(experience) or experience < 0:
                    raise ValueError(f"Row {row_number}: experience_years must be a non-negative number.")
            else:
                experience = None
            source_id = str(row.get("id") or "").strip()
            identity = source_id or hashlib.sha256(f"{name.casefold()}\n{role.casefold()}".encode()).hexdigest()[:16]
            employee_id = f"excel-{identity}"
            if employee_id in employees:
                raise ValueError(f"Row {row_number}: duplicate employee record.")
            employees[employee_id] = {
                "id": employee_id,
                "name": name,
                "role": role,
                "experience_years": experience,
                "skills": [skill.strip() for skill in str(row.get("skills") or "").split(",") if skill.strip()],
            }
        if not employees:
            raise ValueError("The workbook contains no employee records.")
        return employees
    finally:
        workbook.close()